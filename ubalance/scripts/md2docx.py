#!/usr/bin/env python3
"""md → docx 一键转换（pandoc + 统一样式）

用法:
    python3 scripts/md2docx.py <in.md> <out.docx> [--footer "页脚文字"]

样式（与 HUATEX/HUAXING 交付规范一致）:
  · 全文 Times New Roman（含表格、页眉页脚）
  · 表格统一细实线框线（Table Grid）
  · 页脚：文字 + 页码（居中）
  · 中越文混排不换行错乱（保留 pandoc 段落结构，仅改样式）
"""
import argparse
import os
import subprocess
import sys
import tempfile

FONT = 'Times New Roman'


def _rfonts(el):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    return rf


def style_doc(path, footer_text=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document(path)

    # 1) 样式层字体
    for st in doc.styles:
        try:
            _rfonts(st.element)
            st.font.name = FONT
        except Exception:
            pass

    # 2) 逐 run 兜底（pandoc 会显式指定字体）
    def fix_runs(container):
        for r in getattr(container, 'runs', []):
            r.font.name = FONT
            _rfonts(r._element)

    for p in doc.paragraphs:
        fix_runs(p)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    fix_runs(p)

    # 3) 表格框线
    for t in doc.tables:
        try:
            t.style = doc.styles['Table Grid']
        except Exception:
            tblPr = t._tbl.tblPr
            borders = OxmlElement('w:tblBorders')
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                e = OxmlElement('w:' + edge)
                e.set(qn('w:val'), 'single')
                e.set(qn('w:sz'), '4')
                e.set(qn('w:space'), '0')
                e.set(qn('w:color'), '000000')
                borders.append(e)
            tblPr.append(borders)

    # 4) 页脚：文字 + 页码
    for section in doc.sections:
        f = section.footer
        p = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if footer_text:
            run = p.add_run(footer_text + '    ')
            run.font.name = FONT
        run = p.add_run()
        run.font.name = FONT
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'PAGE')
        run._element.addnext(fld)

    doc.save(path)
    return doc


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('src')
    ap.add_argument('dst')
    ap.add_argument('--footer', default=None)
    a = ap.parse_args()

    tmp = os.path.join(tempfile.gettempdir(), '_md2docx_raw.docx')
    subprocess.run(['pandoc', a.src, '-o', tmp, '--from=gfm', '--standalone'], check=True)
    doc = style_doc(tmp, a.footer)
    os.makedirs(os.path.dirname(os.path.abspath(a.dst)), exist_ok=True)
    doc.save(a.dst)
    print('✅ %s → %s（段落 %d / 表格 %d）' % (
        a.src, a.dst, len(doc.paragraphs), len(doc.tables)))


if __name__ == '__main__':
    main()

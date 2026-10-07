#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""recon_ledger.py — 发票级对账台账（D3/F2，新增 2026-10-07）

输入：`payprep to-bank / gen_payment_recon_v2` 产出的同名 sidecar `<定稿表>.recon.json`
输出：台账 xlsx（两个 sheet）
  · 台账        一行 = 一张发票（卷号/发票号/序列号/日期/税号/发票金额/本次付款/备注/
                识别来源/号码来源/是否OCR/正式电子发票确认/OA/收款方/文件）
  · 差异清单    卷级（读值 vs 明细 vs 本次付款 + 口径）＋ 发票级（金额不符/带备注）
                ＋ 待人工核（未匹配明细行 / 非 VND 行）

用法:
  python3 scripts/recon_ledger.py --json "out/Payment request to Bank-20261005-r2.xlsx.recon.json" \
      --out "out/Payment核销-台账-20261005.xlsx"
"""
import argparse
import datetime
import json
import os

HDR_FONT = 'Times New Roman'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--json', required=True, help='定稿表同名 .recon.json')
    ap.add_argument('--out', required=True, help='台账 xlsx 输出路径')
    a = ap.parse_args()

    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    with open(a.json, encoding='utf-8') as f:
        side = json.load(f)

    wb = openpyxl.Workbook()
    thin = Side(style='thin', color='000000')
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _style(ws, ncol, widths=None):
        for c in range(1, ncol + 1):
            for r in range(1, ws.max_row + 1):
                cell = ws.cell(r, c)
                f = cell.font
                cell.font = Font(name=HDR_FONT, size=f.size or 11, bold=(r == 1))
                cell.border = bd
                cell.alignment = Alignment(vertical='center', wrap_text=(c == ncol))
            ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 16)

    # ── Sheet1 台账 ──
    ws = wb.active
    ws.title = '台账'
    cols = ['卷号', '发票号', '序列号', '发票日期', '税号', '发票金额', '本次付款',
            '备注', '识别来源', '号码来源', '是否OCR', '正式电子发票', '请款单号', '收款方', '源文件']
    ws.append(cols)
    for r in side.get('ledger', []):
        ws.append([r.get('folder'), r.get('no'), r.get('serial'), r.get('date'), r.get('tax'),
                   r.get('full_amount'), r.get('pay'), r.get('note'), r.get('src'),
                   r.get('num_src'), 'Y' if r.get('ocr') else '',
                   'Y' if r.get('official') else 'N', r.get('oa'), r.get('vendor'), r.get('file')])
    _style(ws, len(cols), {8: 46, 15: 34, 6: 18, 7: 18})
    print('✅ 台账 %d 行（= 发票行数）' % max(0, ws.max_row - 1))

    # ── Sheet2 差异清单 ──
    ws2 = wb.create_sheet('差异清单')
    ws2.append(['层级', '卷号', '请款单号', '收款方', '发票金额', '本次付款', '明细金额',
                '差额', '口径', '备注/说明'])
    n_diff = 0
    for g in side.get('groups', []):
        det = int(g.get('detail_amount') or 0)
        pay = int(g.get('pay_sum') or 0)
        read = int(g.get('read_sum') or 0)
        if det != pay or read != det or g.get('method') not in ('exact',):
            ws2.append(['卷级', g.get('folder'), g.get('oa'), g.get('payee'), read, pay, det,
                        pay - det, g.get('method'), g.get('note') or ''])
            n_diff += 1
    for r in side.get('ledger', []):
        if int(r.get('full_amount') or 0) != int(r.get('pay') or 0) or r.get('note'):
            ws2.append(['发票级', r.get('folder'), r.get('oa'), r.get('vendor'),
                        r.get('full_amount'), r.get('pay'), '', int(r.get('pay') or 0) - int(r.get('full_amount') or 0),
                        r.get('method'), r.get('note') or ''])
            n_diff += 1
    for r in side.get('unmatched', []):
        ws2.append(['待人工核', r.get('folder'), r.get('oa'), r.get('payee'), '', '',
                    r.get('amount'), '', 'no-invoice', r.get('note') or '明细行未匹配到附件卷'])
        n_diff += 1
    for r in side.get('nonvnd', []):
        ws2.append(['非VND', r.get('folder'), r.get('oa'), r.get('payee'), '', '',
                    r.get('amount'), '', r.get('currency'), r.get('note') or '非 VND 行，另表出'])
        n_diff += 1
    _style(ws2, 10, {9: 14, 10: 50})
    print('✅ 差异清单 %d 行' % n_diff)

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    print('✅ 输出 → %s' % a.out)


if __name__ == '__main__':
    main()

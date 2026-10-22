#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""HMA24011 支付调整函 — 中越双语版生成（xlsx + docx）"""
import shutil, os
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

SRC = "/Users/zhaoyuzhao/.openclaw/media/inbound/HMA24011支付调整函---3b62acd5-b9ab-44dc-8132-bd73e15093ed.xlsx"
OUTDIR = "/Users/zhaoyuzhao/.openclaw/workspace-balance/output"
os.makedirs(OUTDIR, exist_ok=True)
OUT_XLSX = os.path.join(OUTDIR, "HMA24011支付调整函_中越双语.xlsx")
shutil.copyfile(SRC, OUT_XLSX)

wb = openpyxl.load_workbook(OUT_XLSX)
ws = wb["调整通知函"]

ws["A4"] = "调整通知函 / THÔNG BÁO ĐIỀU CHỈNH THANH TOÁN"
ws["A5"] = "TO／Kính gửi：HUA FENG GROUP INC."
ws["A6"] = "事由／Nội dung"

ws["A7"] = (
    "2026年8月7日支付给贵司、银行回单号1032526H07523115的货款，"
    "原付款单据中发票号HMA24011项下的支付金额USD 1,440,660.42填写有误，"
    "实际支付明细应按下表拆分。\n"
    "Ngày 07/08/2026, khoản thanh toán cho Quý Công ty theo chứng từ ngân hàng "
    "số 1032526H07523115, số tiền USD 1.440.660,42 được ghi nhận dưới số hóa đơn "
    "HMA24011 trên chứng từ thanh toán ban đầu là không chính xác. "
    "Số tiền thực tế được phân bổ theo bảng chi tiết dưới đây."
)
ws["A7"].alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

# 原文件含外部链接公式（指向另一工作簿），改为静态数值
ws["B11"] = 1887200.00
ws["C11"] = 1440660.42
ws["B15"] = 1887200.00
ws["C15"] = 587200.00
ws["B16"] = 541495.03
ws["C16"] = 432820.42
ws["B17"] = 420640.00
ws["C17"] = 420640.00
ws["C18"] = 1440660.42
for rng in ("B11", "C11", "B15", "C15", "B16", "C16", "B17", "C17", "B18", "C18"):
    ws[rng].number_format = "#,##0.00"

ws["A9"] = "调整前／TRƯỚC ĐIỀU CHỈNH"
ws["A10"] = "Số Hoá đơn/ Hợp đồng\n发票号／合同号"
ws["B10"] = "发票金额\nGiá trị hóa đơn (USD)"
ws["C10"] = "支付金额\nSố tiền thanh toán (USD)"

ws["A13"] = "调整后／SAU ĐIỀU CHỈNH"
ws["A14"] = "Số Hoá đơn/ Hợp đồng\n发票号／合同号"
ws["B14"] = "发票金额\nGiá trị hóa đơn (USD)"
ws["C14"] = "支付金额\nSố tiền thanh toán (USD)"

ws["A18"] = "合计／Tổng cộng"

ws["A21"] = "确认盖章／XÁC NHẬN VÀ ĐÓNG DẤU"
ws["A23"] = "HUATEX VIET NAM COMPANY LIMITED\n（盖章／Đóng dấu & Ký tên）"
ws["C23"] = "HUA FENG GROUP INC.\n（盖章／Đóng dấu & Ký tên）"

# 编号/日期占位行（插在原 A6 之前不好动，放在 A3）
ws["A3"] = "编号／Số: HMA24011-ADJ-01            日期／Ngày: ______/______/2026"

# 格式统一：表头/内容换行居中
for rng in ("A10", "B10", "C10", "A14", "B14", "C14"):
    ws[rng].alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
for rng in ("A11", "B11", "C11", "A15", "B15", "C15", "A16", "B16", "C16",
            "A17", "B17", "C17"):
    ws[rng].alignment = Alignment(vertical="center", horizontal="center")
ws["A23"].alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
ws["C23"].alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

# 行高
ws.row_dimensions[7].height = 60
ws.row_dimensions[8].height = 30
ws.row_dimensions[10].height = 34
ws.row_dimensions[14].height = 34
ws.row_dimensions[23].height = 44
for col, w in (("A", 34), ("B", 20), ("C", 22), ("D", 10), ("E", 10)):
    ws.column_dimensions[col].width = w
ws.column_dimensions["A"].width = 38

wb.save(OUT_XLSX)
print("xlsx ->", OUT_XLSX)

# ---------------- DOCX ----------------
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.5); s.right_margin = Cm(2.5)

def p(text, size=11, bold=False, align="left", space=6, italic=False):
    par = doc.add_paragraph()
    par.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                     "right": WD_ALIGN_PARAGRAPH.RIGHT, "just": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    par.paragraph_format.space_after = Pt(space)
    run = par.add_run(text)
    run.bold = bold; run.italic = italic; run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    return par

p("HUATEX VIET NAM COMPANY LIMITED", 13, True, "center", 2)
p("ROAD NO. 6, NHON TRACH 6 INDUSTRIAL PARK, PHUOC AN COMMUNE, DONG NAI CITY, VIETNAM", 9, False, "center", 10)
p("Số／编号: HMA24011-ADJ-01", 10, False, "left", 2)
p("Ngày／日期: ______/______/2026", 10, False, "left", 12)

p("调整通知函", 15, True, "center", 2)
p("THÔNG BÁO ĐIỀU CHỈNH THANH TOÁN", 13, True, "center", 14)

p("Kính gửi／致：HUA FENG GROUP INC.", 11, True, "left", 10)

p("事由／Nội dung:", 11, True, "left", 4)
p("2026年8月7日支付给贵司、银行回单号 1032526H07523115 的货款，原付款单据中发票号 HMA24011 "
  "项下的支付金额 USD 1,440,660.42 填写有误，实际支付明细应按下表拆分。",
  11, False, "just", 4)
p("Ngày 07/08/2026, khoản thanh toán cho Quý Công ty theo chứng từ ngân hàng số 1032526H07523115, "
  "số tiền USD 1.440.660,42 được ghi nhận dưới số hóa đơn HMA24011 trên chứng từ thanh toán ban đầu "
  "là không chính xác. Số tiền thực tế được phân bổ theo bảng chi tiết dưới đây.",
  11, False, "just", 12)

def make_table(title_cn, title_vn, rows, total=None, start_row=0):
    p(f"{title_cn}／{title_vn}", 11, True, "left", 4)
    n = len(rows) + 1 + (1 if total else 0)
    t = doc.add_table(rows=n, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ["Số Hoá đơn/ Hợp đồng\n发票号／合同号", "发票金额／Giá trị hóa đơn (USD)",
           "支付金额／Số tiền thanh toán (USD)"]
    for i, h in enumerate(hdr):
        c = t.cell(0, i); c.text = ""
        par = c.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(h); r.bold = True; r.font.size = Pt(10); r.font.name = "Times New Roman"
    for ri, row in enumerate(rows, start=1):
        for ci, v in enumerate(row):
            c = t.cell(ri, ci); c.text = ""
            par = c.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci else WD_ALIGN_PARAGRAPH.LEFT
            r = par.add_run(v); r.font.size = Pt(10.5); r.font.name = "Times New Roman"
    if total:
        i = n - 1
        t.cell(i, 0).text = ""
        par = t.cell(i, 0).paragraphs[0]
        r = par.add_run("合计／Tổng cộng"); r.bold = True; r.font.size = Pt(10.5); r.font.name = "Times New Roman"
        t.cell(i, 1).text = ""
        for ci, v in ((1, ""), (2, total)):
            c = t.cell(i, ci); c.text = ""
            par = c.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(v); r.bold = True; r.font.size = Pt(10.5); r.font.name = "Times New Roman"
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    return t

make_table("调整前", "TRƯỚC ĐIỀU CHỈNH",
           [("HMA24011", "1,887,200.00", "1,440,660.42")])
make_table("调整后", "SAU ĐIỀU CHỈNH",
           [("HMA24011", "1,887,200.00", "587,200.00"),
            ("HMVFZ240042", "541,495.03", "432,820.42"),
            ("HMM24040", "420,640.00", "420,640.00")],
           total="1,440,660.42")

p("说明／Ghi chú:", 11, True, "left", 4)
p("1. 本次调整仅涉及发票级应收/应付明细的重分类，付款总额 USD 1,440,660.42 及已发生的资金往来不变，"
  "不涉及发票重开、价格折让及税务申报金额变更。", 10.5, False, "just", 2)
p("   Việc điều chỉnh này chỉ liên quan đến việc phân bổ lại chi tiết công nợ theo từng hóa đơn; "
  "tổng số tiền thanh toán USD 1.440.660,42 và dòng tiền thực tế không thay đổi, không phát sinh việc "
  "xuất lại hóa đơn, không điều chỉnh giá và không thay đổi số liệu kê khai thuế.", 10.5, False, "just", 6)
p("2. 双方应在 2026 年 8 月所属会计期间的应收/应付明细账中按上表作对应调整，并留存本函作为调账依据。",
  10.5, False, "just", 2)
p("   Hai Bên thực hiện điều chỉnh tương ứng trên sổ chi tiết công nợ phải thu/phải trả thuộc kỳ kế toán "
  "tháng 08/2026 theo bảng trên và lưu giữ Thông báo này làm căn cứ điều chỉnh sổ sách.", 10.5, False, "just", 14)

p("确认盖章／XÁC NHẬN VÀ ĐÓNG DẤU", 11, True, "left", 20)
t = doc.add_table(rows=2, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = True
t.cell(0, 0).text = ""
par = t.cell(0, 0).paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = par.add_run("HUATEX VIET NAM COMPANY LIMITED\n（盖章／Đóng dấu & Ký tên）"); r.bold = True; r.font.size = Pt(10.5)
t.cell(0, 1).text = ""
par = t.cell(0, 1).paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = par.add_run("HUA FENG GROUP INC.\n（盖章／Đóng dấu & Ký tên）"); r.bold = True; r.font.size = Pt(10.5)
for i, (label,) in enumerate([("Ngày／日期: ____/____/2026",), ("Ngày／日期: ____/____/2026",)]):
    c = t.cell(1, i); c.text = ""
    par = c.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run(label); r.font.size = Pt(10)
doc.add_paragraph()

OUT_DOCX = os.path.join(OUTDIR, "HMA24011支付调整函_中越双语.docx")
doc.save(OUT_DOCX)
print("docx ->", OUT_DOCX)

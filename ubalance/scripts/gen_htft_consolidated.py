# -*- coding: utf-8 -*-
"""
Generate consolidated bilingual (EN/VI) Loan Agreement
HT/FT 012025 - Huatex Vietnam (Lender) -> Viet Nam Future Textile (Borrower)
Merges: Original Loan Agreement (02 Jan 2025) + Amendment (730 days / revolving)
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = "/Users/zhaoyuzhao/.openclaw/workspace-balance/output"
os.makedirs(OUT, exist_ok=True)
PATH = os.path.join(OUT, "Hop_dong_vay_HT-FT_012025_Hop_nhat_EN-VI.docx")

doc = Document()

# ---------- base styles ----------
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
pf = st.paragraph_format
pf.space_after = Pt(4)
pf.line_spacing = 1.15

for s in doc.sections:
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(2.0)


def para(text="", bold=False, italic=False, align=None, size=11, space_before=0,
         space_after=4, indent=None, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if color:
        r.font.color.rgb = color
    return p


def pair(en, vi, bold_en=False, space_before=6, indent=None):
    """EN paragraph (normal) followed by VI paragraph (italic)."""
    para(en, bold=bold_en, space_before=space_before, indent=indent)
    para(vi, italic=True, space_before=0, space_after=6, indent=indent)


# =========================================================
# TITLE
# =========================================================
para("LOAN AGREEMENT", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=0)
para("HỢP ĐỒNG VAY", bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=2)
para("(Consolidated version incorporating the Amendment to the Loan Agreement)",
     italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=0)
para("(Bản hợp nhất bao gồm Phụ lục sửa đổi Hợp đồng vay)",
     italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=10)

t = doc.add_table(rows=2, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = True
rows = [
    ("No. / Số:", "HT/FT 012025"),
    ("Original execution date / Ngày ký ban đầu:", "02 January 2025  /  ngày 02 tháng 01 năm 2025"),
]
for i, (k, v) in enumerate(rows):
    c0, c1 = t.rows[i].cells
    c0.text = ""
    c1.text = ""
    r0 = c0.paragraphs[0].add_run(k); r0.bold = True; r0.font.size = Pt(10.5)
    r1 = c1.paragraphs[0].add_run(v); r1.font.size = Pt(10.5)

para("", space_after=2)
para("Place of execution / Nơi ký kết: Dong Nai, Vietnam  /  Đồng Nai, Việt Nam",
     italic=True, size=10.5, space_after=10)

# =========================================================
# PARTIES
# =========================================================
para("THIS AGREEMENT is entered into between:", bold=True, space_before=6, space_after=0)
para("Hợp đồng này được ký kết giữa:", bold=True, italic=True, space_after=8)

pair(
    "1. HUATEX VIETNAM COMPANY LIMITED, a company duly incorporated and operating under the "
    "laws of Vietnam, with Enterprise Registration Certificate No. 3603695597, having its "
    "registered office at No. 6 Street, Nhon Trach VI Industrial Park, Phuoc An Commune, "
    "Dong Nai Province, Vietnam, duly represented by Mr. CHEN, HUI (hereinafter referred to as "
    "the \u201cLender\u201d);",
    "1. CÔNG TY TNHH HUATEX VIỆT NAM, một công ty được thành lập và hoạt động theo pháp luật "
    "Việt Nam, Giấy chứng nhận đăng ký doanh nghiệp số 3603695597, địa chỉ trụ sở: Đường số 6, "
    "Khu công nghiệp Nhơn Trạch VI, Xã Phước An, Tỉnh Đồng Nai, Việt Nam, được đại diện hợp pháp "
    "bởi ông CHEN, HUI (sau đây gọi là \u201cBên Cho Vay\u201d);",
    space_before=4,
)

pair(
    "2. VIET NAM FUTURE TEXTILE COMPANY LIMITED, a company duly incorporated and operating "
    "under the laws of Vietnam, with Enterprise Registration Certificate No. 3901352789, having "
    "its registered office at Lot B11.3, Street C4, Thanh Thanh Cong Industrial Park, Trang Bang "
    "Ward, Tay Ninh Province, Vietnam, duly represented by Mr. HUANG, CHAO GUI (hereinafter "
    "referred to as the \u201cBorrower\u201d).",
    "2. CÔNG TY TNHH FUTURE TEXTILE VIỆT NAM, một công ty được thành lập và hoạt động theo pháp "
    "luật Việt Nam, Giấy chứng nhận đăng ký doanh nghiệp số 3901352789, địa chỉ trụ sở: Lô B11.3, "
    "đường C4, Khu công nghiệp Thành Thành Công, Phường Trảng Bàng, Tỉnh Tây Ninh, Việt Nam, được "
    "đại diện hợp pháp bởi ông HUANG, CHAO GUI (sau đây gọi là \u201cBên Vay\u201d).",
    space_before=4,
)

pair("(each referred to as a \u201cParty\u201d, and jointly as the \u201cParties\u201d).",
     "(mỗi bên sau đây gọi là \u201cBên\u201d và cùng nhau gọi là \u201cCác Bên\u201d).",
     space_before=2)

# =========================================================
# RECITALS
# =========================================================
para("RECITALS", bold=True, space_before=8, space_after=0)
para("LỜI MỞ ĐẦU", bold=True, italic=True, space_after=6)

pair(
    "A. WHEREAS, the Parties entered into Loan Agreement No. HT/FT 012025 dated 02 January 2025 "
    "(the \u201cOriginal Agreement\u201d), and subsequently entered into an Amendment to the Loan "
    "Agreement (the \u201cAmendment\u201d) which amended Clause 3.2(a) of the Original Agreement and "
    "added a new Clause 1.3 thereto; and",
    "A. XÉT RẰNG, Các Bên đã ký kết Hợp đồng Vay số HT/FT 012025 ngày 02 tháng 01 năm 2025 "
    "(\u201cHợp Đồng Gốc\u201d), và sau đó đã ký kết Phụ lục sửa đổi Hợp đồng Vay "
    "(\u201cVăn Bản Sửa Đổi\u201d) để sửa đổi Điều khoản 3.2(a) của Hợp Đồng Gốc và bổ sung "
    "Điều khoản 1.3 mới; và",
    space_before=4,
)

pair(
    "B. WHEREAS, the Parties wish to consolidate the Original Agreement and the Amendment into a "
    "single agreement reflecting all terms as currently in force; and",
    "B. XÉT RẰNG, Các Bên mong muốn hợp nhất Hợp Đồng Gốc và Văn Bản Sửa Đổi thành một văn bản "
    "duy nhất phản ánh toàn bộ các điều khoản đang có hiệu lực; và",
    space_before=4,
)

pair(
    "C. WHEREAS, the Lender has agreed to lend to the Borrower, and the Borrower has agreed to "
    "borrow, the Loan (as defined below) on the terms and conditions set out herein.",
    "C. XÉT RẰNG, Bên Cho Vay đồng ý cho Bên Vay vay và Bên Vay đồng ý vay Khoản Vay (như được "
    "định nghĩa dưới đây) theo các điều khoản và điều kiện quy định trong Hợp đồng này.",
    space_before=4,
)

para("NOW, THEREFORE, the Parties agree as follows:", bold=True, space_before=8, space_after=0)
para("VÌ VẬY, Các Bên đồng ý như sau:", bold=True, italic=True, space_after=8)

# =========================================================
# 0. CONSOLIDATION STATUS
# =========================================================
para("0. STATUS OF THIS AGREEMENT", bold=True, space_before=6, space_after=0)
para("0. TÌNH TRẠNG CỦA HỢP ĐỒNG NÀY", bold=True, italic=True, space_after=6)

pair(
    "0.1. This Agreement consolidates, restates and supersedes the Original Agreement and the "
    "Amendment, which shall from the Consolidation Date no longer have separate legal effect. "
    "Any reference to the Original Agreement or the Amendment in any other document shall, from "
    "the Consolidation Date, be read as a reference to this Agreement.",
    "0.1. Hợp đồng này hợp nhất, trình bày lại và thay thế Hợp Đồng Gốc và Văn Bản Sửa Đổi; "
    "kể từ Ngày Hợp Nhất, các văn bản đó không còn hiệu lực riêng lẻ. Mọi dẫn chiếu đến Hợp Đồng "
    "Gốc hoặc Văn Bản Sửa Đổi trong bất kỳ tài liệu nào khác, kể từ Ngày Hợp Nhất, được hiểu là "
    "dẫn chiếu đến Hợp đồng này.",
    space_before=4,
)

pair(
    "0.2. For the avoidance of doubt, the Parties confirm that: (i) this Agreement does not "
    "constitute a new loan and does not discharge, novate or otherwise extinguish any obligation "
    "of the Borrower outstanding as at the Consolidation Date; (ii) all Drawdowns made prior to "
    "the Consolidation Date remain outstanding and continue to be governed by the terms of this "
    "Agreement; and (iii) the aggregate principal amount outstanding under this Agreement shall "
    "not exceed the Loan amount set out in Clause 1.1.",
    "0.2. Để tránh hiểu nhầm, Các Bên xác nhận rằng: (i) Hợp đồng này không cấu thành một khoản "
    "vay mới và không làm giải phóng, thay thế nghĩa vụ hay chấm dứt bất kỳ nghĩa vụ nào của "
    "Bên Vay còn tồn đọng tại Ngày Hợp Nhất; (ii) toàn bộ các Lần Giải Ngân thực hiện trước Ngày "
    "Hợp Nhất vẫn còn dư nợ và tiếp tục được điều chỉnh bởi các điều khoản của Hợp đồng này; và "
    "(iii) tổng số tiền gốc còn dư nợ theo Hợp đồng này không vượt quá hạn mức quy định tại "
    "Điều 1.1.",
    space_before=4,
)

pair(
    "0.3. The Parties confirm that the Original Agreement was executed on 02 January 2025, and "
    "that any reference in the Vietnamese text of the Original Agreement to \u201cngày 02 tháng 10 "
    "năm 2025\u201d is a clerical error and shall be read as \u201cngày 02 tháng 01 năm 2025\u201d.",
    "0.3. Các Bên xác nhận rằng Hợp Đồng Gốc đã được ký kết vào ngày 02 tháng 01 năm 2025, và "
    "mọi dẫn chiếu trong bản tiếng Việt của Hợp Đồng Gốc đến \u201cngày 02 tháng 10 năm 2025\u201d "
    "là lỗi đánh máy và phải được hiểu là \u201cngày 02 tháng 01 năm 2025\u201d.",
    space_before=4,
)

pair(
    "0.4. This Agreement takes effect on the date of its execution by both Parties (the "
    "\u201cConsolidation Date\u201d).",
    "0.4. Hợp đồng này có hiệu lực kể từ ngày được cả hai Bên ký kết (\u201cNgày Hợp Nhất\u201d).",
    space_before=4,
)

# =========================================================
# 1. THE LOAN
# =========================================================
para("1. THE LOAN", bold=True, space_before=10, space_after=0)
para("1. KHOẢN VAY", bold=True, italic=True, space_after=6)

pair(
    "1.1. Subject to the terms and conditions set forth herein, the Lender shall lend to the "
    "Borrower the aggregate original principal amount of up to One Hundred and Thirty-Two Billion "
    "Vietnamese Dong (VND 132,000,000,000) (the \u201cLoan\u201d).",
    "1.1. Theo các điều khoản và điều kiện được quy định trong Hợp đồng này, Bên Cho Vay sẽ cho "
    "Bên Vay vay tổng số tiền gốc ban đầu lên đến Một Trăm Ba Mươi Hai Tỷ Đồng Việt Nam "
    "(132.000.000.000 VND) (sau đây gọi là \u201cKhoản Vay\u201d).",
    space_before=4,
)

pair(
    "1.2. The purpose of the Loan is the replenishment for all operations of the Borrower in "
    "compliance with applicable laws.",
    "1.2. Mục đích của Khoản Vay là bổ sung nguồn vốn cho toàn bộ hoạt động của Bên Vay phù hợp "
    "với quy định pháp luật.",
    space_before=4,
)

pair(
    "1.3. Revolving Credit Facility. The Loan constitutes a revolving credit facility. Within the "
    "limit set out in Clause 1.1 and subject to the terms of this Agreement, the Borrower may draw "
    "down, repay and re-draw the Loan by way of one or more Drawdowns, provided that at no time "
    "shall the aggregate principal amount outstanding under this Agreement exceed the amount set "
    "out in Clause 1.1.",
    "1.3. Hạn mức tín dụng quay vòng. Khoản Vay cấu thành hạn mức tín dụng quay vòng. Trong phạm "
    "vi hạn mức quy định tại Điều 1.1 và theo các điều kiện của Hợp đồng này, Bên Vay được quyền "
    "giải ngân, hoàn trả và tiếp tục giải ngân lại thành một hoặc nhiều Lần Giải Ngân; với điều "
    "kiện tại bất kỳ thời điểm nào, tổng số tiền gốc còn dư nợ theo Hợp đồng này không được vượt "
    "quá hạn mức quy định tại Điều 1.1.",
    space_before=4,
)

# =========================================================
# 2. DRAWDOWN
# =========================================================
para("2. DRAWDOWN", bold=True, space_before=10, space_after=0)
para("2. GIẢI NGÂN", bold=True, italic=True, space_after=6)

pair(
    "2.1. The Loan shall be drawn down in one or more installments (each installment a "
    "\u201cDrawdown\u201d).",
    "2.1. Khoản Vay sẽ được giải ngân thành một hoặc nhiều đợt (mỗi đợt gọi là \u201cLần Giải "
    "Ngân\u201d).",
    space_before=4,
)

pair(
    "2.2. The Parties agree that a Drawdown shall be considered disbursed in full upon the "
    "transfer of the amount of that Drawdown from the Lender\u2019s bank account to the "
    "Borrower\u2019s Account. Each Drawdown shall have its own \u201cDrawdown Date\u201d, being the "
    "date on which the relevant Drawdown is credited to the Borrower\u2019s Account. The date on "
    "which the first Drawdown is credited to the Borrower\u2019s Account shall be the \u201cFirst "
    "Drawdown Date\u201d.",
    "2.2. Các Bên đồng ý rằng một Lần Giải Ngân được coi là đã giải ngân đầy đủ khi số tiền của "
    "Lần Giải Ngân đó được chuyển từ tài khoản ngân hàng của Bên Cho Vay sang tài khoản của Bên "
    "Vay. Mỗi Lần Giải Ngân có \u201cNgày Giải Ngân\u201d riêng, là ngày Lần Giải Ngân đó được ghi "
    "có vào tài khoản của Bên Vay. Ngày Lần Giải Ngân đầu tiên được ghi có vào tài khoản của Bên "
    "Vay được gọi là \u201cNgày Giải Ngân Đầu Tiên\u201d.",
    space_before=4,
)

# =========================================================
# 3. REPAYMENT
# =========================================================
para("3. REPAYMENT", bold=True, space_before=10, space_after=0)
para("3. HOÀN TRẢ", bold=True, italic=True, space_after=6)

pair(
    "3.1. The Borrower shall repay each Drawdown on the respective Repayment Date (as defined "
    "below).",
    "3.1. Bên Vay sẽ hoàn trả mỗi Lần Giải Ngân vào Ngày Hoàn Trả tương ứng (được quy định dưới "
    "đây).",
    space_before=4,
)

pair(
    "3.2. \u201cRepayment Date\u201d for each respective Drawdown means the date which:",
    "3.2. \u201cNgày Hoàn Trả\u201d cho mỗi Lần Giải Ngân tương ứng được xác định là:",
    space_before=4,
)

pair(
    "(a) falls 730 days from the First Drawdown Date, being the final date on which the Loan must "
    "be repaid in full (the \u201cFinal Maturity Date\u201d); provided that the Borrower shall "
    "have the right to repay the Loan, in whole or in part, at any time prior to the Final "
    "Maturity Date without any penalty;",
    "(a) là ngày cách Ngày Giải Ngân Đầu Tiên 730 ngày, là ngày cuối cùng mà Khoản Vay phải được "
    "hoàn trả toàn bộ (\u201cNgày Đáo Hạn Cuối Cùng\u201d); với điều kiện Bên Vay có quyền hoàn "
    "trả Khoản Vay, toàn bộ hoặc một phần, vào bất kỳ thời điểm nào trước Ngày Đáo Hạn Cuối Cùng "
    "mà không phải chịu bất kỳ khoản phạt nào;",
    space_before=4,
    indent=0.6,
)

pair(
    "(b) in case the Repayment Date on which the repayment of the relevant Drawdown must be "
    "performed falls on a non-business day in Vietnam, the Repayment Date shall be the next "
    "following business day.",
    "(b) trong trường hợp Ngày Hoàn Trả mà việc hoàn trả Lần Giải Ngân tương ứng phải được thực "
    "hiện rơi vào ngày không làm việc tại Việt Nam, thì Ngày Hoàn Trả sẽ là ngày làm việc kế tiếp.",
    space_before=2,
    indent=0.6,
)

pair(
    "3.3. Each Drawdown shall be repaid on the respective Repayment Date by wire transfer "
    "according to the payment details of the Lender as specified in Section 7 (the "
    "\u201cLender\u2019s Account\u201d).",
    "3.3. Mỗi Lần Giải Ngân sẽ được hoàn trả vào Ngày Hoàn Trả tương ứng bằng hình thức chuyển "
    "khoản theo thông tin thanh toán của Bên Cho Vay được quy định tại Mục 7 "
    "(\u201cTài Khoản Bên Cho Vay\u201d).",
    space_before=4,
)

# =========================================================
# 4. INTEREST
# =========================================================
para("4. INTEREST", bold=True, space_before=10, space_after=0)
para("4. LÃI SUẤT", bold=True, italic=True, space_after=6)

pair(
    "4.1. Interest shall be calculated and accrue on the basis of the actual number of days "
    "elapsed and a 365-day year on the outstanding total amount of all Drawdowns disbursed under "
    "the Loan, starting from the respective Drawdown Date of each Drawdown.",
    "4.1. Lãi suất sẽ được tính và tích lũy trên cơ sở số ngày thực tế đã trôi qua và một năm "
    "365 ngày, trên tổng số tiền còn dư nợ của tất cả các Lần Giải Ngân đã giải ngân theo Khoản "
    "Vay, tính từ Ngày Giải Ngân của từng Lần Giải Ngân.",
    space_before=4,
)

pair(
    "4.2. The interest on all Drawdowns disbursed under the Loan shall be due and payable on the "
    "Repayment Date of the last outstanding Drawdown.",
    "4.2. Lãi suất đối với tất cả các Lần Giải Ngân đã giải ngân theo Khoản Vay sẽ đến hạn và "
    "phải được thanh toán vào Ngày Hoàn Trả của Lần Giải Ngân còn dư nợ cuối cùng.",
    space_before=4,
)

pair(
    "4.3. The interest rate applicable to each Drawdown shall be 4% per annum.",
    "4.3. Lãi suất áp dụng đối với mỗi Lần Giải Ngân là 4%/năm.",
    space_before=4,
)

# =========================================================
# 5. MISCELLANEOUS
# =========================================================
para("5. MISCELLANEOUS", bold=True, space_before=10, space_after=0)
para("5. CÁC ĐIỀU KHOẢN KHÁC", bold=True, italic=True, space_after=6)

pair(
    "5.1. Any change and/or amendment to this Agreement shall be made by the Parties through the "
    "execution of an additional agreement signed by authorised persons.",
    "5.1. Mọi thay đổi và/hoặc sửa đổi đối với Hợp đồng này phải được Các Bên thực hiện thông "
    "qua việc ký kết một văn bản bổ sung do các đại diện có thẩm quyền ký.",
    space_before=4,
)

pair(
    "5.2. All payments to be made by the Borrower shall be made to the Lender in immediately "
    "available and freely transferable funds, for full value, on the due date, to such account as "
    "the Lender may specify, subject to applicable laws and regulations.",
    "5.2. Tất cả các khoản thanh toán của Bên Vay phải được thực hiện cho Bên Cho Vay bằng nguồn "
    "tiền sẵn có và chuyển nhượng tự do, đủ giá trị, vào ngày đến hạn, đến tài khoản do Bên Cho "
    "Vay chỉ định, phù hợp với quy định pháp luật hiện hành.",
    space_before=4,
)

pair(
    "5.3. Payment of all expenses related to the Loan, including correspondent banks\u2019 fees "
    "and commissions, shall be borne by the beneficiary of the relevant payment.",
    "5.3. Việc thanh toán tất cả các chi phí liên quan đến Khoản Vay, bao gồm phí và hoa hồng "
    "của các ngân hàng đại lý, sẽ do bên nhận khoản thanh toán tương ứng chịu.",
    space_before=4,
)

pair(
    "5.4. Each of the Lender and the Borrower shall bear and pay its own costs in connection with "
    "the preparation and execution of this Agreement. Each Party shall pay its own stamp, "
    "documentary, registration or other like duties or taxes imposed on or in connection with "
    "this Agreement or the Loan, and shall indemnify the other Party against any liability arising "
    "by reason of any delay or omission to pay such duties or taxes.",
    "5.4. Mỗi Bên Cho Vay và Bên Vay tự chịu và thanh toán các chi phí của mình liên quan đến "
    "việc chuẩn bị và ký kết Hợp đồng này. Mỗi Bên tự thanh toán các loại lệ phí, thuế trước bạ, "
    "thuế tài liệu, thuế đăng ký hoặc các loại thuế tương tự khác liên quan đến Hợp đồng này hoặc "
    "Khoản Vay, và bồi thường cho Bên kia đối với bất kỳ nghĩa vụ nào phát sinh do chậm trễ hoặc "
    "thiếu sót trong việc thanh toán các khoản đó.",
    space_before=4,
)

pair(
    "5.5. No failure or delay on the part of the Lender in exercising any power, right or remedy "
    "under this Agreement shall operate as a waiver of such power, right or remedy, nor shall any "
    "single or partial exercise thereof preclude any other or further exercise thereof or the "
    "exercise of any other power, right or remedy.",
    "5.5. Việc Bên Cho Vay không thực hiện hoặc chậm thực hiện bất kỳ quyền hạn, quyền hoặc biện "
    "pháp khắc phục nào theo Hợp đồng này sẽ không được coi là từ bỏ quyền hạn, quyền hoặc biện "
    "pháp khắc phục đó; và việc thực hiện một phần hoặc một lần quyền hạn, quyền hoặc biện pháp "
    "khắc phục sẽ không ngăn cản việc thực hiện bất kỳ quyền hạn, quyền hoặc biện pháp khắc phục "
    "khác hoặc việc thực hiện thêm.",
    space_before=4,
)

pair(
    "5.6. This Agreement takes effect upon its signing by both Parties and remains valid until the "
    "Parties have fully performed their respective obligations hereunder.",
    "5.6. Hợp đồng này có hiệu lực ngay khi được cả hai Bên ký kết và vẫn còn hiệu lực cho đến "
    "khi Các Bên hoàn thành đầy đủ các nghĩa vụ của mình theo Hợp đồng này.",
    space_before=4,
)

# =========================================================
# 6. GOVERNING LAW
# =========================================================
para("6. GOVERNING LAW AND JURISDICTION", bold=True, space_before=10, space_after=0)
para("6. LUẬT ÁP DỤNG VÀ CƠ QUAN GIẢI QUYẾT TRANH CHẤP", bold=True, italic=True, space_after=6)

pair(
    "6.1. This Agreement shall be governed by and construed in accordance with the laws of "
    "Vietnam.",
    "6.1. Hợp đồng này được điều chỉnh và giải thích theo pháp luật Việt Nam.",
    space_before=4,
)

pair(
    "6.2. Any dispute arising out of or in connection with this Agreement (including any dispute "
    "regarding its existence, validity or termination) shall be submitted to the competent "
    "People\u2019s Court of Vietnam for final resolution.",
    "6.2. Mọi tranh chấp phát sinh từ hoặc liên quan đến Hợp đồng này (bao gồm tranh chấp về sự "
    "tồn tại, tính hiệu lực hoặc việc chấm dứt Hợp đồng) sẽ được đưa ra Tòa án nhân dân có thẩm "
    "quyền của Việt Nam để giải quyết cuối cùng.",
    space_before=4,
)

# =========================================================
# 7. PAYMENT DETAILS
# =========================================================
para("7. PAYMENT DETAILS", bold=True, space_before=10, space_after=0)
para("7. THÔNG TIN THANH TOÁN", bold=True, italic=True, space_after=6)

pair(
    "Payments under this Agreement shall be made to and from the bank accounts of the Parties as "
    "set out below. Each Party shall notify the other Party in writing of any change to its bank "
    "details at least 05 business days prior to the effective date of such change.",
    "Các khoản thanh toán theo Hợp đồng này được thực hiện đến và từ các tài khoản ngân hàng của "
    "Các Bên như dưới đây. Mỗi Bên phải thông báo bằng văn bản cho Bên kia về mọi thay đổi thông "
    "tin ngân hàng của mình ít nhất 05 ngày làm việc trước ngày thay đổi có hiệu lực.",
    space_before=2,
)


def bank_block(title_en, title_vi, lines):
    para(title_en, bold=True, space_before=8, space_after=0)
    para(title_vi, bold=True, italic=True, space_after=3)
    for ln in lines:
        para(ln, size=10.5, space_after=1, indent=0.5)


bank_block(
    "The Borrower: VIET NAM FUTURE TEXTILE COMPANY LIMITED",
    "Bên Vay: CÔNG TY TNHH FUTURE TEXTILE VIỆT NAM",
    [
        "Legal Address / Địa chỉ trụ sở: Lot B11.3, Street C4, Thanh Thanh Cong Industrial Park, "
        "Trang Bang Ward, Tay Ninh Province, Vietnam",
        "Địa chỉ: Lô B11.3, đường C4, Khu công nghiệp Thành Thành Công, Phường Trảng Bàng, "
        "Tỉnh Tây Ninh, Việt Nam",
        "Account No. (VND) / Số tài khoản (VND): 1863636868 \u2013 Military Commercial Joint Stock "
        "Bank (MB Bank) \u2013 Tam Hiep Branch",
        "Account No. (VND) / Số tài khoản (VND): 111002998317 \u2013 Vietnam Joint Stock Commercial "
        "Bank for Industry and Trade (Vietinbank) \u2013 Nhon Trach Branch",
    ],
)

bank_block(
    "The Lender: HUATEX VIETNAM COMPANY LIMITED",
    "Bên Cho Vay: CÔNG TY TNHH HUATEX VIỆT NAM",
    [
        "Legal Address / Địa chỉ trụ sở: No. 6 Street, Nhon Trach VI Industrial Park, Phuoc An "
        "Commune, Dong Nai Province, Vietnam",
        "Địa chỉ: Đường số 6, Khu công nghiệp Nhơn Trạch VI, Xã Phước An, Tỉnh Đồng Nai, Việt Nam",
        "Account No. (VND) / Số tài khoản (VND): 111002923670 \u2013 Vietnam Joint Stock Commercial "
        "Bank for Industry and Trade (Vietinbank)",
        "Account No. (VND) / Số tài khoản (VND): 5656568568 \u2013 Military Commercial Joint Stock "
        "Bank (MB Bank)",
    ],
)

# =========================================================
# 8. GENERAL TERMS
# =========================================================
para("8. GENERAL TERMS", bold=True, space_before=10, space_after=0)
para("8. ĐIỀU KHOẢN CHUNG", bold=True, italic=True, space_after=6)

pair(
    "8.1. This Agreement is made in 04 original copies in both the English and Vietnamese "
    "languages. Each Party shall keep 02 original copies, all of which shall have equal legal "
    "effect.",
    "8.1. Hợp đồng này được lập thành 04 bản gốc bằng cả tiếng Anh và tiếng Việt. Mỗi Bên giữ 02 "
    "bản gốc, tất cả các bản đều có giá trị pháp lý như nhau.",
    space_before=4,
)

pair(
    "8.2. This Agreement is made in English and Vietnamese. In the event of any inconsistency "
    "between the two language versions, the Vietnamese version shall prevail.",
    "8.2. Hợp đồng này được lập bằng tiếng Anh và tiếng Việt. Trong trường hợp có bất kỳ sự không "
    "nhất quán nào giữa hai bản ngôn ngữ, bản tiếng Việt sẽ được ưu tiên áp dụng.",
    space_before=4,
)

pair(
    "8.3. In witness whereof, the Parties have signed this Agreement on the date written below.",
    "8.3. Để ghi nhận sự thỏa thuận, Các Bên đã ký kết Hợp đồng này vào ngày ghi dưới đây.",
    space_before=4,
)

# =========================================================
# SIGNATURES
# =========================================================
para("", space_after=10)

sig = doc.add_table(rows=1, cols=2)
sig.autofit = True
left, right = sig.rows[0].cells

def fill(cell, en, vi, company, rep, role_en, role_vi):
    p = cell.paragraphs[0]
    r = p.add_run(en); r.bold = True; r.font.size = Pt(10.5)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(vi); r2.bold = True; r2.italic = True; r2.font.size = Pt(10.5)
    for txt, it in [(company, False), (rep, False)]:
        pp = cell.add_paragraph()
        rr = pp.add_run(txt); rr.font.size = Pt(10.5); rr.italic = it
    cell.add_paragraph()
    cell.add_paragraph()
    pp = cell.add_paragraph()
    rr = pp.add_run("_____________________________"); rr.font.size = Pt(10.5)
    pp = cell.add_paragraph()
    rr = pp.add_run(role_en); rr.font.size = Pt(10.5); rr.bold = True
    pp = cell.add_paragraph()
    rr = pp.add_run(role_vi); rr.font.size = Pt(10.5); rr.italic = True

fill(left, "For and on behalf of The Borrower", "Thay mặt và đại diện cho Bên Vay",
     "VIET NAM FUTURE TEXTILE COMPANY LIMITED", "(Signature and seal / Ký tên và đóng dấu)",
     "HUANG, CHAO GUI", "Legal Representative / Người đại diện theo pháp luật")

fill(right, "For and on behalf of The Lender", "Thay mặt và đại diện cho Bên Cho Vay",
     "HUATEX VIETNAM COMPANY LIMITED", "(Signature and seal / Ký tên và đóng dấu)",
     "CHEN, HUI", "Legal Representative / Người đại diện theo pháp luật")

para("", space_after=8)
para("Date / Ngày: .............. / .............. / 20......",
     size=10.5, space_before=6)

doc.save(PATH)
print("SAVED:", PATH)
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))

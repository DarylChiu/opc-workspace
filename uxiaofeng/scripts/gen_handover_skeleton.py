#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""接收容器：① 工作交接清单（Daryl Amend 版结构，待填）"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, BLUE, LIGHT, GREY = "1F3864", "2E5C8A", "DCE6F1", "F7F7F7"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = openpyxl.Workbook()


def title_block(ws, title, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 6


def build(name, title, headers, widths, nrows=25, validations=None):
    ws = wb.create_sheet(name)
    title_block(ws, title, len(headers))
    r0 = 3
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=r0, column=i, value=h)
        c.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[r0].height = 32
    for r in range(r0 + 1, r0 + 1 + nrows):
        ws.row_dimensions[r].height = 30
        for ci in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=ci)
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True,
                                       horizontal="center" if ci == 1 else "left",
                                       indent=0 if ci == 1 else 1)
            cell.border = BORDER
            if (r - r0) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=GREY)
    ws.freeze_panes = f"A{r0 + 1}"
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + nrows}"
    if validations:
        for col_idx, options in validations:
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            L = get_column_letter(col_idx)
            dv.add(f"{L}{r0+1}:{L}{r0+nrows}")
    return ws


ws = wb.active
ws.title = "① 交接概要"
title_block(ws, "工作交接清单", 2)
rows = [
    ("交接人", "赵毓钊"),
    ("接手人", ""),
    ("交接人岗位", "财务高级专员"),
    ("所属部门", "海外财务课"),
    ("直属上级", "陈剑锋"),
    ("预计离职/调岗日期", "2026-10-30"),
    ("交接启动日期", "2026-10-26"),
    ("交接完成日期", "2026-10-29"),
    ("交接方式", "□ 面谈  □ 文档  □ 系统演示  □ 陪同实操"),
    ("交接范围概述", ""),
    ("遗留风险提示", ""),
]
r = 3
for k, v in rows:
    c1 = ws.cell(row=r, column=1, value=k)
    c1.font = Font(name="微软雅黑", size=10, bold=True)
    c1.fill = PatternFill("solid", fgColor=LIGHT)
    c1.alignment = Alignment(vertical="center", indent=1)
    c1.border = BORDER
    c2 = ws.cell(row=r, column=2, value=v)
    c2.font = Font(name="微软雅黑", size=10)
    c2.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    c2.border = BORDER
    ws.row_dimensions[r].height = 28
    r += 1
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 80

build("② 工作职责清单", "工作职责清单",
      ["序号", "职责模块", "具体工作内容", "执行频率", "关键时间节点/截止日", "接手人", "当前状态", "备注"],
      [6, 20, 50, 12, 26, 10, 11, 34],
      validations=[(4, ["每日", "每周", "每月", "每季", "每年", "不定期", "项目制"]),
                   (7, ["已交接", "进行中", "未开始"])])

build("③ 待办与进行中事项", "待办与进行中事项",
      ["序号", "事项名称", "背景/来龙去脉", "当前进度", "下一步动作", "截止日期", "相关方/对接人", "风险等级", "资料存放位置"],
      [6, 24, 52, 14, 46, 12, 20, 9, 26],
      validations=[(8, ["高", "中", "低"])])

build("④ 文档与资料索引", "文档与资料索引",
      ["序号", "资料名称", "类型", "存储路径 / 链接（共享盘）", "负责人", "更新频率", "是否已确认可访问"],
      [6, 36, 12, 42, 10, 12, 16],
      validations=[(3, ["制度文件", "合同协议", "财务台账", "报表模板", "操作手册", "邮件记录", "其他"]),
                   (7, ["是", "否"])])

build("⑤ 系统与账号权限", "系统与账号权限 (Systems & Access)",
      ["序号", "系统/平台名称", "用途", "账号持有人", "权限级别", "密码", "权限回收/移交方式", "已处理"],
      [6, 30, 26, 12, 12, 24, 40, 10],
      validations=[(5, ["管理员", "普通用户", "只读", "审批权限"]), (8, ["是", "否", "不适用"])])

build("⑥ 关键联系人与对接方", "关键联系人与对接方",
      ["序号", "姓名", "单位/部门", "职务", "负责事项", "联系方式", "注意事项"],
      [6, 14, 26, 16, 34, 18, 42])

build("⑦ 风险与注意事项", "风险与注意事项",
      ["序号", "风险点", "潜在影响", "应对方式"],
      [6, 44, 52, 56])

out = "/Users/zhaoyuzhao/.openclaw/xiaofeng_workspace/output/工作交接清单_赵毓钊_正式版.xlsx"
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)
print("saved:", out)
